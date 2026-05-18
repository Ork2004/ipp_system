import re
from datetime import date
from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from backend.app.config import GENERATED_DIR
from backend.app.database import get_connection


KEY_TEACHER = "ФИО ППС"
KEY_POSITION = "Должность"
KEY_KIND = "вид занятии"
KEY_DISCIPLINE = "Дисциплина"
KEY_TOTAL = "Итого"
KEY_L = "Л"
KEY_SPZ = "СПЗ"
KEY_LZ = "ЛЗ"
KEY_SRSP = "СРСП"
KEY_RK = "РК 1,2"
KEY_EXAM = "экзамены"

SEMESTER_KEY_RE = re.compile(r"^\s*(\d+)\s*семестр\s*$", re.IGNORECASE)


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _normalize_text_lower(value: Any) -> str:
    return _normalize_text(value).lower()


def _to_num(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ".").strip())
    except Exception:
        return 0.0


def _format_num(value: Any) -> str:
    value = round(_to_num(value), 2)
    if abs(value - int(value)) < 1e-9:
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def _safe_filename(value: str) -> str:
    return re.sub(r"[^\w]+", "_", str(value or ""), flags=re.UNICODE).strip("_")


def _short_teacher_name(full_name: str) -> str:
    parts = [part for part in _normalize_text(full_name).split(" ") if part]
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    initials = "".join(f"{part[0]}." for part in parts[1:] if part)
    return f"{parts[0]} {initials}".strip()


def _date_ru(value: date) -> str:
    months = {
        1: "января",
        2: "февраля",
        3: "марта",
        4: "апреля",
        5: "мая",
        6: "июня",
        7: "июля",
        8: "августа",
        9: "сентября",
        10: "октября",
        11: "ноября",
        12: "декабря",
    }
    return f"{value.day:02d} {months[value.month]} {value.year} г."


def _blank_teacher_profile(teacher_name: str, department_name: str = "") -> dict:
    return {
        "id": None,
        "full_name": teacher_name,
        "department_id": None,
        "department": department_name,
        "faculty": "",
        "position": "",
        "academic_degree": "",
        "academic_rank": "",
        "staff_type": "",
    }


def _load_department_name(cur, department_id: Optional[int]) -> str:
    if not department_id:
        return ""
    cur.execute("SELECT name FROM departments WHERE id=%s;", (department_id,))
    row = cur.fetchone()
    return row[0] if row else ""


def _load_teacher_profile(cur, teacher_name: str, department_id: Optional[int]) -> dict:
    params: list[Any] = [_normalize_text(teacher_name)]
    q = """
        SELECT
            t.id,
            t.full_name,
            t.department_id,
            d.name,
            t.faculty,
            t.position,
            t.academic_degree,
            t.academic_rank,
            t.staff_type
        FROM teachers t
        LEFT JOIN departments d ON d.id = t.department_id
        WHERE lower(t.full_name) = lower(%s)
    """
    if department_id:
        q += " AND t.department_id = %s"
        params.append(department_id)
    q += " ORDER BY t.id LIMIT 1;"

    cur.execute(q, params)
    row = cur.fetchone()
    if not row:
        return _blank_teacher_profile(teacher_name, _load_department_name(cur, department_id))

    return {
        "id": row[0],
        "full_name": row[1] or teacher_name,
        "department_id": row[2],
        "department": row[3] or "",
        "faculty": row[4] or "",
        "position": row[5] or "",
        "academic_degree": row[6] or "",
        "academic_rank": row[7] or "",
        "staff_type": row[8] or "",
    }


def _load_teacher_rows(cur, teacher_name: str, academic_year: str, department_id: Optional[int]) -> tuple[list[dict], str]:
    params: list[Any] = [academic_year]
    q = """
        SELECT er.row_data, d.name
        FROM excel_rows er
        JOIN excel_templates et ON et.id = er.template_id
        LEFT JOIN departments d ON d.id = et.department_id
        WHERE et.status = 'parsed'
          AND et.academic_year = %s
    """
    if department_id:
        q += " AND et.department_id = %s"
        params.append(department_id)
    q += " ORDER BY er.row_number;"

    cur.execute(q, params)
    rows = []
    department_name = ""
    target_name = _normalize_text_lower(teacher_name)

    for row_data, dept_name in cur.fetchall() or []:
        if not isinstance(row_data, dict):
            continue
        row_teacher_name = _normalize_text_lower(row_data.get(KEY_TEACHER))
        if row_teacher_name != target_name:
            continue
        rows.append(row_data)
        if dept_name and not department_name:
            department_name = dept_name

    return rows, department_name


def _active_semesters(row_data: dict) -> list[int]:
    semesters: list[int] = []
    for key, value in (row_data or {}).items():
        match = SEMESTER_KEY_RE.match(str(key))
        if not match:
            continue
        if _normalize_text(value) == "" or abs(_to_num(value)) < 1e-9:
            continue
        sem_num = int(match.group(1))
        if sem_num not in semesters:
            semesters.append(sem_num)
    return sorted(semesters)


def _empty_totals() -> dict:
    return {
        "total": 0.0,
        "l": 0.0,
        "spz": 0.0,
        "lz": 0.0,
        "srsp": 0.0,
        "rk": 0.0,
        "exam": 0.0,
        "mooc": 0.0,
        "diploma": 0.0,
        "research": 0.0,
    }


def _is_mooc_row(row_data: dict) -> bool:
    kind = _normalize_text_lower(row_data.get(KEY_KIND))
    discipline = _normalize_text_lower(row_data.get(KEY_DISCIPLINE))
    return "моок" in kind or "mooc" in kind or "coursera" in discipline


def _is_diploma_row(row_data: dict) -> bool:
    text = _normalize_text_lower(
        " ".join([
            str(row_data.get(KEY_KIND) or ""),
            str(row_data.get(KEY_DISCIPLINE) or ""),
        ])
    )
    return any(token in text for token in ("диплом", "диссертац", "дп", "мд"))


def _is_research_row(row_data: dict) -> bool:
    text = _normalize_text_lower(
        " ".join([
            str(row_data.get(KEY_KIND) or ""),
            str(row_data.get(KEY_DISCIPLINE) or ""),
        ])
    )
    return any(token in text for token in ("нир", "науч", "research"))


def _build_report_context(rows: list[dict]) -> dict:
    semesters = {1: {}, 2: {}}
    totals_by_semester = {1: _empty_totals(), 2: _empty_totals()}
    annual = _empty_totals()
    disciplines = set()

    for row_data in rows:
        discipline = _normalize_text(row_data.get(KEY_DISCIPLINE)) or "Без названия"
        total = _to_num(row_data.get(KEY_TOTAL))
        active_semesters = _active_semesters(row_data)
        if active_semesters:
            share = 1.0 / len(active_semesters)
        else:
            share = 1.0

        values = {
            "total": total,
            "l": _to_num(row_data.get(KEY_L)),
            "spz": _to_num(row_data.get(KEY_SPZ)),
            "lz": _to_num(row_data.get(KEY_LZ)),
            "srsp": _to_num(row_data.get(KEY_SRSP)),
            "rk": _to_num(row_data.get(KEY_RK)),
            "exam": _to_num(row_data.get(KEY_EXAM)),
            "mooc": total if _is_mooc_row(row_data) else 0.0,
            "diploma": total if _is_diploma_row(row_data) else 0.0,
            "research": total if _is_research_row(row_data) else 0.0,
        }

        for key, value in values.items():
            annual[key] += value

        if discipline:
            disciplines.add(discipline)

        for sem_num in active_semesters:
            if sem_num not in semesters:
                continue
            semesters[sem_num][discipline] = semesters[sem_num].get(discipline, 0.0) + total * share
            for key, value in values.items():
                totals_by_semester[sem_num][key] += value * share

    return {
        "semester_disciplines": semesters,
        "totals_by_semester": totals_by_semester,
        "annual": annual,
        "disciplines": sorted(disciplines),
    }


def _set_default_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def _add_paragraph(doc: Document, text: str = "", *, bold: bool = False, center: bool = False):
    paragraph = doc.add_paragraph()
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    return paragraph


def _add_section(doc: Document, title: str) -> None:
    paragraph = _add_paragraph(doc, title, bold=True)
    paragraph.paragraph_format.space_before = Pt(8)


def _add_table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = "" if value is None else str(value)


def _discipline_rows(discipline_map: dict[str, float]) -> list[list[Any]]:
    rows = []
    for idx, (discipline, hours) in enumerate(
        sorted(discipline_map.items(), key=lambda item: item[0].lower()),
        start=1,
    ):
        rows.append([idx, discipline, f"{_format_num(hours)} ч."])
    if not rows:
        rows.append(["", "Нет данных по семестру", ""])
    return rows


def _join_limited(items: list[str], limit: int = 6) -> str:
    filtered = [item for item in items if item]
    if not filtered:
        return "нет данных"
    visible = filtered[:limit]
    suffix = "" if len(filtered) <= limit else f" и еще {len(filtered) - limit}"
    return ", ".join(visible) + suffix


def _add_summary_table(doc: Document, context: dict) -> None:
    sem_totals = context["totals_by_semester"]
    rows = []
    for sem_num, label in ((1, "Осенний семестр"), (2, "Весенний семестр")):
        totals = sem_totals.get(sem_num) or _empty_totals()
        control = totals["rk"] + totals["exam"]
        rows.append([
            label,
            _format_num(totals["srsp"]),
            _format_num(control),
            _format_num(totals["srsp"] + control),
        ])
    _add_table(
        doc,
        ["Учебная нагрузка", "СРС/СРСП", "Рубежный контроль и экзамены", "Всего часов"],
        rows,
    )


def _add_final_table(doc: Document, context: dict) -> None:
    annual = context["annual"]
    auditory = annual["l"] + annual["spz"] + annual["lz"]
    extra = max(annual["total"] - auditory, 0.0)
    scientific = annual["research"] + annual["mooc"]
    rows = [
        [1, "Учебная работа / Аудиторная", _format_num(auditory), _format_num(auditory)],
        [1, "Учебная работа / Внеаудиторная", _format_num(extra), _format_num(extra)],
        [2, "Учебно-методическая работа", "0", "0"],
        [3, "Научная работа", _format_num(scientific), _format_num(scientific)],
        [4, "Организационно-методическая работа", "0", "0"],
        [5, "Воспитательная, профориентационная и общественная работа", "0", "0"],
        ["", "Итого", _format_num(annual["total"] + scientific), _format_num(annual["total"] + scientific)],
    ]
    _add_table(doc, ["№", "Виды работ", "План", "Выполнение"], rows)


def _build_output_path(teacher_name: str, academic_year: str) -> Path:
    safe_teacher = _safe_filename(teacher_name) or "teacher"
    safe_year = _safe_filename(academic_year) or "year"
    return (Path(GENERATED_DIR) / f"Report_{safe_teacher}_{safe_year}.docx").resolve()


def _build_summary_context(rows: list[dict]) -> dict:
    teachers: dict[str, float] = {}
    annual = _empty_totals()
    semester_totals = {1: _empty_totals(), 2: _empty_totals()}

    for row_data in rows:
        teacher_name = _normalize_text(row_data.get(KEY_TEACHER)) or "Без имени"
        total = _to_num(row_data.get(KEY_TOTAL))
        active_semesters = _active_semesters(row_data)
        share = 1.0 / len(active_semesters) if active_semesters else 1.0

        values = {
            "total": total,
            "l": _to_num(row_data.get(KEY_L)),
            "spz": _to_num(row_data.get(KEY_SPZ)),
            "lz": _to_num(row_data.get(KEY_LZ)),
            "srsp": _to_num(row_data.get(KEY_SRSP)),
            "rk": _to_num(row_data.get(KEY_RK)),
            "exam": _to_num(row_data.get(KEY_EXAM)),
            "mooc": total if _is_mooc_row(row_data) else 0.0,
            "diploma": total if _is_diploma_row(row_data) else 0.0,
            "research": total if _is_research_row(row_data) else 0.0,
        }

        for key, value in values.items():
            annual[key] += value

        teachers[teacher_name] = teachers.get(teacher_name, 0.0) + total

        for sem_num in active_semesters:
            if sem_num not in semester_totals:
                continue
            for key, value in values.items():
                semester_totals[sem_num][key] += value * share

    top_teachers = sorted(
        [{"name": name, "total": total} for name, total in teachers.items()],
        key=lambda x: x["total"],
        reverse=True,
    )

    return {
        "annual": annual,
        "totals_by_semester": semester_totals,
        "semester_disciplines": {1: {}, 2: {}},
        "disciplines": [],
        "teachers_count": len(teachers),
        "top_teachers": top_teachers,
    }


def _build_summary_output_path(academic_year: str, department_id: Optional[int] = None) -> Path:
    dept_part = f"_dept{department_id}" if department_id else "_all"
    safe_year = _safe_filename(academic_year) or "year"
    return (Path(GENERATED_DIR) / f"Summary_Report{dept_part}_{safe_year}.docx").resolve()


def generate_summary_report_docx(
    *,
    academic_year: str,
    department_id: Optional[int] = None,
) -> str:
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            params: list[Any] = [academic_year]
            q = """
                SELECT er.row_data, d.name
                FROM excel_rows er
                JOIN excel_templates et ON et.id = er.template_id
                LEFT JOIN departments d ON d.id = et.department_id
                WHERE et.status = 'parsed'
                  AND et.academic_year = %s
            """
            if department_id:
                q += " AND et.department_id = %s"
                params.append(department_id)
            cur.execute(q, params)

            rows: list[dict] = []
            dept_names: set[str] = set()
            for row_data, dept_name in cur.fetchall() or []:
                if isinstance(row_data, dict):
                    rows.append(row_data)
                if dept_name:
                    dept_names.add(dept_name)

            if department_id:
                dept_display = _load_department_name(cur, department_id) or ", ".join(sorted(dept_names)) or "кафедры"
            else:
                dept_display = ", ".join(sorted(dept_names)) or "всех кафедр"

        context = _build_summary_context(rows)
        annual = context["annual"]
        sem1 = context["totals_by_semester"][1]
        sem2 = context["totals_by_semester"][2]

        doc = Document()
        _set_default_styles(doc)

        _add_paragraph(doc, f"Сводный отчёт по нагрузке {dept_display}", bold=True, center=True)
        _add_paragraph(doc, f"за {academic_year} учебный год", center=True)
        _add_paragraph(doc)
        _add_paragraph(doc, f"Количество преподавателей: {context['teachers_count']}")
        _add_paragraph(doc, f"Кафедра: {dept_display}")
        _add_paragraph(doc)
        _add_paragraph(doc, f"За {academic_year} учебный год преподавателями выполнена следующая нагрузка:")

        _add_section(doc, "Учебная работа")
        _add_paragraph(doc, f"а) Выполнена нагрузка в осеннем семестре: {_format_num(sem1['total'])} часов.")
        _add_paragraph(doc, f"б) Выполнена нагрузка в весеннем семестре: {_format_num(sem2['total'])} часов.")
        _add_paragraph(doc, f"в) Перевыполнение и недовыполнение нагрузки в связи с праздниками, командировкой, больничным: не указано.")
        _add_paragraph(doc, f"г) Руководство дипломными проектами и магистерскими диссертациями: {_format_num(annual['diploma'])} ч.")

        _add_paragraph(doc)
        _add_paragraph(doc, "Нагрузка преподавателей:", bold=True)
        top_rows = [
            [i + 1, t["name"], f"{_format_num(t['total'])} ч."]
            for i, t in enumerate(context["top_teachers"][:10])
        ] or [["", "Нет данных", ""]]
        _add_table(doc, ["№", "ФИО преподавателя", "Объём нагрузки"], top_rows)

        _add_section(doc, "Учебная работа по проверке успеваемости обучающихся")
        _add_summary_table(doc, context)

        _add_section(doc, "Учебно-методическая работа")
        _add_paragraph(doc, "Сведения об использовании инновационных технологий обучения: заполняется кафедрой.")

        _add_section(doc, "Научно-исследовательская работа")
        _add_paragraph(doc, f"а) МООК и онлайн-курсы: {_format_num(annual['mooc'])} ч.")
        _add_paragraph(doc, f"б) Научно-исследовательская работа: {_format_num(annual['research'])} ч.")
        _add_paragraph(doc, "в) Список публикаций: заполняется кафедрой.")

        _add_section(doc, "Организационно-методическая работа")
        _add_paragraph(doc, "Организационно-методическая работа: заполняется кафедрой.")

        _add_section(doc, "Воспитательная, профориентационная, общественная работа")
        _add_paragraph(doc, "Воспитательная работа: заполняется кафедрой.")

        _add_section(doc, f"Итоги выполнения нагрузки преподавателей за {academic_year} учебный год")
        _add_final_table(doc, context)
        _add_paragraph(doc)
        _add_paragraph(doc, _date_ru(date.today()))

        output_path = _build_summary_output_path(academic_year, department_id)
        doc.save(output_path)
        return str(output_path)
    finally:
        conn.close()


def generate_teacher_report_docx(
    *,
    teacher_name: str,
    academic_year: str,
    department_id: Optional[int] = None,
) -> str:
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            rows, excel_department_name = _load_teacher_rows(cur, teacher_name, academic_year, department_id)
            teacher = _load_teacher_profile(cur, teacher_name, department_id)
            if not teacher.get("department") and excel_department_name:
                teacher["department"] = excel_department_name

        context = _build_report_context(rows)

        doc = Document()
        _set_default_styles(doc)

        full_name = teacher.get("full_name") or teacher_name
        short_name = _short_teacher_name(full_name)
        department = teacher.get("department") or "кафедры"
        degree_rank = ", ".join(
            item for item in (teacher.get("academic_degree"), teacher.get("academic_rank")) if item
        )

        _add_paragraph(doc, f"Отчет преподавателя кафедры {department} {short_name}", bold=True, center=True)
        _add_paragraph(doc, f"о выполнении индивидуального плана за {academic_year} учебный год", center=True)
        _add_paragraph(doc)
        _add_paragraph(doc, f"ФИО: {full_name}")
        _add_paragraph(doc, f"Должность: {teacher.get('position') or 'не указано'}")
        _add_paragraph(doc, f"Степень, звание: {degree_rank or 'не указано'}")
        _add_paragraph(doc, f"Штатный, совместитель: {teacher.get('staff_type') or 'не указано'}")
        _add_paragraph(doc)
        _add_paragraph(doc, f"За {academic_year} учебный год мной проделана следующая работа:")

        _add_section(doc, "Учебная работа")
        sem1_total = context["totals_by_semester"][1]["total"]
        sem2_total = context["totals_by_semester"][2]["total"]
        _add_paragraph(doc, f"а) Выполнена нагрузка в осеннем семестре: {_format_num(sem1_total)} часов.")
        _add_paragraph(doc, f"Читаемые дисциплины за осенний семестр {academic_year} учебного года.")
        _add_table(doc, ["№", "Преподаваемый предмет, дисциплина", "Объем"], _discipline_rows(context["semester_disciplines"][1]))
        _add_paragraph(doc)
        _add_paragraph(doc, f"б) Выполнена нагрузка в весеннем семестре: {_format_num(sem2_total)} часов.")
        _add_paragraph(doc, f"Читаемые дисциплины за весенний семестр {academic_year} учебного года.")
        _add_table(doc, ["№", "Преподаваемый предмет, дисциплина", "Объем"], _discipline_rows(context["semester_disciplines"][2]))
        _add_paragraph(doc)
        _add_paragraph(doc, "в) Перевыполнение и недовыполнение нагрузки в связи с праздниками, командировкой, больничным: не указано.")
        _add_paragraph(doc, f"г) Руководство дипломными проектами и магистерскими диссертациями: {_format_num(context['annual']['diploma'])} ч. по данным нагрузки.")
        _add_paragraph(doc, f"д) Элективные и авторские курсы: {_join_limited(context['disciplines'])}.")

        _add_section(doc, "Учебная работа по проверке успеваемости обучающихся")
        _add_summary_table(doc, context)

        _add_section(doc, "Учебно-методическая работа")
        _add_paragraph(doc, "а) Семинары: не указано.")
        _add_paragraph(doc, "б) Сведения об использовании инновационных технологий обучения: заполняется преподавателем.")

        _add_section(doc, "Научно-исследовательская работа")
        _add_paragraph(doc, "а) Список публикаций")
        _add_table(
            doc,
            ["№", "Ф.И.О. автор(ы)", "Название публикации", "Выходные данные"],
            [[1, "Заполняется преподавателем", "", ""]],
        )
        _add_paragraph(doc, "б) Участие в научных проектах: заполняется преподавателем.")
        _add_paragraph(doc, "в) Сертификаты, награды: заполняется преподавателем.")

        _add_section(doc, "Организационно-методическая работа")
        _add_paragraph(doc, "Организационно-методическая работа: заполняется преподавателем.")

        _add_section(doc, "Воспитательная, профориентационная, общественная работа")
        _add_paragraph(doc, "Повышение квалификации и общественная работа: заполняется преподавателем.")

        _add_section(doc, f"Итоги выполнения ИП работы преподавателя за {academic_year} учебный год")
        _add_final_table(doc, context)
        _add_paragraph(doc)
        _add_paragraph(doc, short_name)
        _add_paragraph(doc, _date_ru(date.today()))

        output_path = _build_output_path(full_name, academic_year)
        doc.save(output_path)
        return str(output_path)
    finally:
        conn.close()
