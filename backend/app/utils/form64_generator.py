"""
Form 64 — department-wide cumulative report (Сводный отчёт по выполнению
педагогической нагрузки ППС кафедры).

Unlike Form 63 (an XLSX with planned hours per teacher), Form 64 is the
official department DOCX where every teacher occupies three rows — 1 семестр,
2 семестр and Итого за уч. год — and each work category is split into a
План / Выполнение pair.

Strategy: keep the official trilingual header and signature blocks intact by
working off the real form as a template. The first teacher's three data rows
in that template are well-formed (consistent gridSpan/vMerge), so we deep-copy
them as prototypes, drop every existing data row, then clone the prototypes
once per teacher and overwrite the cell text by raw <w:tc> index.

The "Выполнение" (actual) columns are intentionally left blank: the system
tracks only planned hours (the IUP summary's Выполнение columns are empty in
practice), so the department fills actuals in by hand.
"""

import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# Header rows (0-based) of the form's main table; data rows start after this.
DATA_START_ROW_INDEX = 7
PROTOTYPE_ROW_INDICES = (7, 8, 9)  # sem1, sem2, yearly total of teacher #1

# Within a semester prototype row the first data <w:tc> is at this index;
# within the yearly-total prototype row it starts earlier (the label spans the
# №..Семестр columns).
SEM_DATA_TC_START = 5
TOTAL_DATA_TC_START = 2

# Ordered plan slots. Each slot maps to one (План, Выполнение) cell pair, so
# 9 slots == 18 data cells, matching both prototype layouts.
PLAN_SLOTS = (
    "auditory",
    "extraauditory",
    "methodical",
    "research",
    "organizational_methodical",
    "educational",
    "total",
    "hourly_teaching",
    "hourly_dvr",
)


def _to_num(value) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _fmt(value) -> str:
    """Format an hour value the way the form does (comma decimals, no .0)."""
    num = round(_to_num(value), 2)
    if abs(num - int(num)) < 1e-9:
        return str(int(num))
    return f"{num:.2f}".rstrip("0").rstrip(".").replace(".", ",")


def _slots_from_form63_item(item: dict) -> dict[str, float]:
    """Map one Form 63 (teacher, semester) item to the 9 Form 64 plan slots."""
    educational = (
        _to_num(item.get("educational_hours"))
        + _to_num(item.get("qualification_hours"))
        + _to_num(item.get("social_hours"))
    )
    return {
        "auditory": _to_num(item.get("teaching_auditory_hours")),
        "extraauditory": _to_num(item.get("teaching_extraauditory_hours")),
        "methodical": _to_num(item.get("methodical_hours")),
        "research": _to_num(item.get("research_hours")),
        "organizational_methodical": _to_num(item.get("organizational_methodical_hours")),
        "educational": educational,
        "total": _to_num(item.get("planned_total_hours")),
        "hourly_teaching": _to_num(item.get("hourly_auditory_hours")),
        "hourly_dvr": _to_num(item.get("hourly_extraauditory_hours")),
    }


def _empty_slots() -> dict[str, float]:
    return {slot: 0.0 for slot in PLAN_SLOTS}


def build_form64_rows_from_form63(items: list[dict]) -> list[dict]:
    """
    Group Form 63 preview items (one per teacher+semester) into Form 64 rows:
    one dict per teacher with sem1 / sem2 plan slots and a derived yearly total.
    """
    grouped: dict[tuple, dict] = {}
    order: list[tuple] = []

    for item in items:
        name = item.get("teacher_name")
        position = item.get("position")
        if not name:
            continue
        key = (name, position)
        if key not in grouped:
            grouped[key] = {
                "teacher_name": name,
                "position": position,
                "rate": item.get("rate", ""),
                "payment": item.get("payment", ""),
                "sem1": _empty_slots(),
                "sem2": _empty_slots(),
            }
            order.append(key)

        slots = _slots_from_form63_item(item)
        sem = "sem1" if item.get("semester") == 1 else "sem2"
        for slot, value in slots.items():
            grouped[key][sem][slot] += value

    rows = []
    for key in order:
        teacher = grouped[key]
        year = {
            slot: teacher["sem1"][slot] + teacher["sem2"][slot]
            for slot in PLAN_SLOTS
        }
        teacher["year"] = year
        rows.append(teacher)
    return rows


# ---------- docx cell helpers ----------

def _set_tc_text(tc, text: str):
    """Overwrite a <w:tc>'s text, preserving its paragraph/run formatting."""
    text = "" if text is None else str(text)

    paragraphs = tc.findall(qn("w:p"))
    if not paragraphs:
        p = OxmlElement("w:p")
        tc.append(p)
        paragraphs = [p]

    p = paragraphs[0]
    for extra in paragraphs[1:]:
        tc.remove(extra)

    runs = p.findall(qn("w:r"))
    rpr_template = None
    if runs:
        first_rpr = runs[0].find(qn("w:rPr"))
        if first_rpr is not None:
            rpr_template = copy.deepcopy(first_rpr)
        for r in runs:
            p.remove(r)

    run = OxmlElement("w:r")
    if rpr_template is not None:
        run.append(rpr_template)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    p.append(run)


def _row_tcs(tr):
    return tr.findall(qn("w:tc"))


def _fill_data_cells(tcs, start_index: int, slots: dict[str, float]):
    """Write План into even data cells, leave Выполнение (odd) cells blank."""
    for i, slot in enumerate(PLAN_SLOTS):
        plan_idx = start_index + 2 * i
        actual_idx = plan_idx + 1
        if plan_idx < len(tcs):
            _set_tc_text(tcs[plan_idx], _fmt(slots.get(slot, 0.0)))
        if actual_idx < len(tcs):
            _set_tc_text(tcs[actual_idx], "")


def _set_header_text(tbl, dept_name: str | None, academic_year: str | None):
    """Best-effort fill of the department name (row 3) and year (row 1)."""
    rows = tbl.rows
    if academic_year and len(rows) > 1:
        tcs = _row_tcs(rows[1]._tr)
        if len(tcs) >= 2:
            _set_tc_text(tcs[1], f"{academic_year} оқу жылы/ учебный год/ academic year")
    if dept_name and len(rows) > 3:
        tcs = _row_tcs(rows[3]._tr)
        if len(tcs) >= 2:
            _set_tc_text(tcs[1], dept_name)


def export_form64_docx(
    rows: list[dict],
    template_path: str,
    output_path: str,
    dept_name: str | None = None,
    academic_year: str | None = None,
):
    """
    Fill the Form 64 DOCX template with one teacher block (3 rows) per item in
    `rows` (as produced by build_form64_rows_from_form63).
    """
    doc = Document(template_path)

    # The official file stacks two Form 64 tables (full-time staff, then a
    # part-time roster with a grand-total). The Form 63 data pipeline already
    # returns every teacher of the department in one list, so we fill only the
    # first table (it carries the trilingual title/approval header) and drop
    # any further tables — otherwise their original rows leak into the output.
    for extra in doc.tables[1:]:
        extra._tbl.getparent().remove(extra._tbl)

    tbl = doc.tables[0]
    tbl_el = tbl._tbl

    all_trs = tbl_el.findall(qn("w:tr"))
    if len(all_trs) <= max(PROTOTYPE_ROW_INDICES):
        raise ValueError("Шаблон Формы 64 повреждён: не хватает строк-прототипов")

    proto_sem1 = copy.deepcopy(all_trs[PROTOTYPE_ROW_INDICES[0]])
    proto_sem2 = copy.deepcopy(all_trs[PROTOTYPE_ROW_INDICES[1]])
    proto_total = copy.deepcopy(all_trs[PROTOTYPE_ROW_INDICES[2]])

    # Drop every existing data row, keeping the header rows only.
    for tr in all_trs[DATA_START_ROW_INDEX:]:
        tbl_el.remove(tr)

    _set_header_text(tbl, dept_name, academic_year)

    for no, teacher in enumerate(rows, start=1):
        sem1_tr = copy.deepcopy(proto_sem1)
        sem2_tr = copy.deepcopy(proto_sem2)
        total_tr = copy.deepcopy(proto_total)

        sem1_tcs = _row_tcs(sem1_tr)
        sem2_tcs = _row_tcs(sem2_tr)
        total_tcs = _row_tcs(total_tr)

        # identity columns live on the sem1 row (vMerge-restart cells)
        if len(sem1_tcs) > 4:
            _set_tc_text(sem1_tcs[0], str(no))
            _set_tc_text(sem1_tcs[1], teacher.get("teacher_name") or "")
            _set_tc_text(sem1_tcs[2], teacher.get("rate") or "")
            _set_tc_text(sem1_tcs[3], teacher.get("payment") or "")
            _set_tc_text(sem1_tcs[4], "1 сем.")
        if len(sem2_tcs) > 4:
            _set_tc_text(sem2_tcs[4], "2 сем.")
        if len(total_tcs) > 1:
            _set_tc_text(total_tcs[1], "Итого за уч. год:")

        _fill_data_cells(sem1_tcs, SEM_DATA_TC_START, teacher["sem1"])
        _fill_data_cells(sem2_tcs, SEM_DATA_TC_START, teacher["sem2"])
        _fill_data_cells(total_tcs, TOTAL_DATA_TC_START, teacher["year"])

        tbl_el.append(sem1_tr)
        tbl_el.append(sem2_tr)
        tbl_el.append(total_tr)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
